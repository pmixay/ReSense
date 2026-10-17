#!/usr/bin/env python3
"""Build P2's prerelease draft from the immutable P4 export, without rescoring."""
import argparse
import hashlib
import json
import subprocess
from collections import defaultdict
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from p2_chart_sources import pointer

ROOT = Path(__file__).resolve().parents[1]
P4 = '6d6aaca'
EXPORT = 'docs/evidence/p4_eval_1.1.0_2026-10-16/export/'


def build(output):
    output.mkdir(parents=True, exist_ok=True)
    sha = subprocess.check_output(['git', 'rev-parse', P4], cwd=ROOT, text=True).strip()
    blobs = {}

    def blob(path):
        if path not in blobs:
            blobs[path] = subprocess.check_output(['git', 'show', f'{sha}:{path}'], cwd=ROOT)
        return blobs[path]

    rows = json.loads(blob(EXPORT + 'numbers.json'))['rows']
    register = json.loads(blob(EXPORT + 'case_register.json'))['rows']
    # Independently verify every direct exported number, not just selected chart cells.
    for row in rows:
        if not row['source'] and row['derived']:
            continue  # register totals; independently checked by P4 build_export --check
        data = blob(row['source'])
        if hashlib.sha256(data).hexdigest() != row['source_sha256']:
            raise ValueError(f"source hash mismatch: {row['source']}")
        if not row['derived'] and pointer(json.loads(data), row['pointer']) != row['value']:
            raise ValueError(f"source value mismatch: {row}")
    selected = [r for r in rows if r['version'] == '1.1.0']
    used = []

    def value(stage, case, metric):
        found = [r for r in selected if (r['stage'], r['case_id'], r['metric']) == (stage, case, metric)]
        if len(found) != 1:
            raise ValueError(f'expected exactly one source: {stage}/{case}/{metric}')
        used.append(found[0])
        return found[0]['value']

    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 12})
    charts = []

    def save(fig, name, words):
        fig.tight_layout()
        for ext in ('png', 'svg'):
            fig.savefig(output / f'{name}.{ext}', dpi=160, metadata={'Date': None} if ext == 'svg' else {})
        svg = output / f'{name}.svg'
        svg.write_text('\n'.join(line.rstrip() for line in svg.read_text().splitlines()) + '\n')
        plt.close(fig)
        charts.append({'name': name, 'visible_words_upper_bound': words})

    cases = [r for r in register if r['stage'] == 'X1' and r['version'] == '1.1.0']
    fig, ax = plt.subplots(figsize=(13, 6))
    labels = []
    for i, case in enumerate(cases):
        cid = case['case_id']
        hits = value('X1', cid, 'stop_hit_frames')
        visible = value('X1', cid, 'visible_frames')
        geometry = {'A': 'прямая', 'B': 'двухпутный', 'C': 'затвор', 'D': 'платформа', 'S': 'стоянка', 'E6': 'прямая', 'setO': 'setO'}[case['alias']]
        labels.append(f"{geometry}/{case['object']}  {hits}/{visible}")
        if case['object'] == 'setO':
            ax.text(3, i, 'окна; дальность не определена', va='center', fontsize=10)
            continue
        for metric, color, marker in [('first_stop_m', '#e4000d', 'o'), ('held_from_m', '#172c60', 's')]:
            if metric == 'held_from_m' and not any(r['case_id'] == cid and r['metric'] == metric for r in selected):
                # P4's numeric export omits nulls. Recover the explicit source null;
                # absence alone must never be interpreted as a measured zero/miss.
                source = json.loads(blob(case['evidence']))
                v = source['cases'][case['object']]['1.1.0']['judge_summary']['held_from_X']
                assert v is None
            else:
                v = value('X1', cid, metric)
            if v is None:
                ax.text(170, i, 'нет held', fontsize=10, va='center')
            else:
                ax.scatter(v, i, c=color, marker=marker, s=45)
                ax.annotate(f'{v:.1f}', (v, i), xytext=(3, 7 if marker == 'o' else -14), textcoords='offset points', fontsize=9)
    ax.set_yticks(range(len(labels)), labels)
    ax.invert_yaxis()
    ax.set_ylim(len(labels) - .5, -.6)
    ax.set_xlim(0, 205)
    ax.set_xticks([0, 50, 100, 150, 200])
    ax.set_xlabel('м · красный: первый STOP; синий: held ≥90% по определению P4')
    ax.grid(axis='x', alpha=.2)
    save(fig, 'range-x1', 68)

    variants = ['nominal', 'pitch_p2', 'pitch_m2', 'roll_p2', 'hz5', 'drop30_s1', 'drop30_s2', 'drop30_s3', 'drop30_judge7', 'dedup']
    fig, axes = plt.subplots(1, 2, figsize=(13, 5))
    for recording, label, color in [('squareT_platform_squareT_switch', 'платформа', '#e4000d'), ('roundT_squareT_pressureGate_squareT', 'roundT-squareT', '#172c60')]:
        vals = [value('X3', f'x3/{recording}/{v}', 'false_stop_episodes') for v in variants]
        axes[0].plot(range(10), vals, 'o-', label=label, color=color)
    ratios = []
    for v in variants:
        cid = f'x3/doubleT_obstacle/{v}'
        hits, total = value('X3', cid, 'stop_frames'), value('X3', cid, 'frames')
        ratios.append((hits, total))
    axes[1].bar(range(10), [h / n for h, n in ratios], color='#172c60')
    for i, (h, n) in enumerate(ratios):
        axes[1].text(i, .04, f'{h}/{n}', rotation=90, color='white', ha='center')
    ticks = ['nom', '+p2', '−p2', '+r2', '5Hz', 'd1', 'd2', 'd3', 'd7', 'dedup']
    for ax in axes:
        ax.set_xticks(range(10), ticks, rotation=45)
        ax.grid(axis='y', alpha=.2)
    axes[0].set_title('Ложные STOP: эпизоды')
    axes[0].legend()
    axes[1].set_title('doubleT_obstacle: STOP / кадры')
    axes[1].set_ylim(0, 1)
    save(fig, 'x3-robustness', 63)

    bins = ['30-60m', '60-100m', '100-150m', '150-200m', '200-260m']
    totals = defaultdict(lambda: [0] * len(bins))
    for row in selected:
        if row['stage'] != 'X1' or not row['metric'].startswith('cause.'):
            continue
        cause, band = row['metric'][6:].rsplit('.', 1)
        group = ('нет кластера' if cause == 'no_blob' else 'отброшен' if cause.startswith('drop:')
                 else 'трекер' if cause.startswith('track:') else 'понижен' if cause.startswith('demoted:') else 'прочее')
        totals[group][bins.index(band)] += row['value']
        used.append(row)
    fig, ax = plt.subplots(figsize=(13, 5))
    bottom = [0] * len(bins)
    for label, vals in sorted(totals.items()):
        ax.bar(bins, vals, bottom=bottom, label=label)
        bottom = [a + b for a, b in zip(bottom, vals)]
    ax.set_ylabel('атрибутированные строки без STOP на объекте')
    ax.legend(ncol=5, loc='upper center', bbox_to_anchor=(.5, 1.15))
    save(fig, 'x1-causes', 55)

    # The named original supplies the unchanged organizer/team slides and photos.
    prs = Presentation(ROOT / 'docs/presentation/ReSense_LCT2026.pptx')
    original = (ROOT / 'docs/presentation/ReSense_LCT2026.pptx').read_bytes()

    def text(slide, x, y, w, h, content, size, bold=False):
        box = slide.shapes.add_textbox(x, y, w, h)
        box.text = content
        for p in box.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            p.font.name = 'Arial'
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string('17171C')

    def chart_slide(index, title, chart, caption, notes):
        slide = prs.slides[index] if index is not None else prs.slides.add_slide(prs.slide_layouts[6])
        slide._element.set('showMasterSp', '0')
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        w, h = prs.slide_width, prs.slide_height
        text(slide, int(w*.04), int(h*.03), int(w*.92), int(h*.10), title, 26, True)
        slide.shapes.add_picture(str(output / f'{chart}.png'), int(w*.03), int(h*.17), width=int(w*.94), height=int(h*.65))
        text(slide, int(w*.04), int(h*.84), int(w*.92), int(h*.14), caption, 12)
        slide.notes_slide.notes_text_frame.text = notes

    common = f'P4 {sha}. Node-decoded; sealed prerelease 1.1.0; in-sample. Linux Haswell, not S5. Release J1 and P1 approval pending. Complete register and chart sources accompany this draft. '
    chart_slide(11, 'ДАЛЬНОСТЬ · X1', 'range-x1', 'ЧЕРНОВИК 1.1.0 · синтетика, in-sample · node-path', common + 'All 11 X1 cases shown; setO comprises selected attribution windows, not a range trial. Counts are STOP hits / visible frames. A: roundT_squareT_pressureGate_squareT; B: roundT_doubleT; C: roundT_pressureGate_roundT; D/S: squareT_platform_squareT_switch. Stationary S cases do not prove moving-train range. held is the P4 judge definition, not continuous STOP or guaranteed detection. A_lowbox_rail is a rail-offset probe, not the below-envelope negative control.')
    chart_slide(None, 'X3 · УСТОЙЧИВОСТЬ', 'x3-robustness', 'ЧЕРНОВИК · in-sample · 1.1.0 = v1.0.0; улучшения нет', common + 'All ten variants are shown. p=pitch, r=roll (degrees); d=30% dropout with seeds 1/2/3/7; dedup removes duplicate returns. All 40 cases and per-object metrics remain in chart-input.json. No held-out claim. Pin BLAS for comparison.')
    chart_slide(None, 'X1 · ПРИЧИНЫ ПОЗДНЕГО STOP', 'x1-causes', 'ЧЕРНОВИК · in-sample · одна причина на строку; не recall', common + 'Sum all exported candidate cause rows over all X1 cases by distance band. First cause in code order. no_blob does not distinguish preprocessing removal from DBSCAN noise. Cases differ in geometry/object and duration; totals are not rates. SetO attribution windows are included.')
    # Replace obsolete object slide with explicitly node-path captions, not relabelled gate bars.
    slide = prs.slides[12]
    slide._element.set('showMasterSp', '0')
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    nominal = 'x3/cloud_with_fake_obj/nominal'
    obj = [('big_edge_inside', 'Краевой ящик 2 м'), ('long_low_on_rails', 'Доска на рельсах'), ('big_above', 'Ящик сверху габарита')]
    captions = []
    for key, label in obj:
        v = value('X3', nominal, f'object.{key}.first_stop_m')
        captions.append(f'{label}: первый STOP {v:.1f} м')
    text(slide, Inches(.5), Inches(.3), prs.slide_width-Inches(1), Inches(.8), 'ОБЪЕКТЫ ОРГАНИЗАТОРОВ', 28, True)
    text(slide, Inches(.6), Inches(1.5), prs.slide_width-Inches(1.2), Inches(3), '\n\n'.join(captions), 26)
    text(slide, Inches(.6), prs.slide_height-Inches(1.2), prs.slide_width-Inches(1.2), Inches(1), 'ЧЕРНОВИК 1.1.0 · node-path · синтетика, in-sample\nОкруглённо: 28 м / 89 м / 111,4 м. Не гарантированная дальность.', 16)
    slide.notes_slide.notes_text_frame.text = common + 'Corrects older 28.09 values, not a detector-improvement claim. Full per-object failures, counts and perturbations in the export.'
    # Every slide is unmistakably draft, including retained historical measurements.
    for slide in prs.slides:
        text(slide, int(prs.slide_width*.85), 0, int(prs.slide_width*.14), Inches(.22), 'ЧЕРНОВИК/16.10', 8)
    deck = output / 'ReSense_LCT2026_DRAFT.pptx'
    prs.save(deck)
    payload = {'p4_commit': sha, 'numbers': rows, 'case_register': register, 'used_rows': used,
               'caption_policy': common, 'charts': charts}
    (output / 'chart-input.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2)+'\n')
    for path, data in blobs.items():
        target = output / 'sources' / path
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(data)
    # Conservative budget includes raster chart labels, which ordinary PPTX audits cannot see.
    counts = [sum(len(s.text.split()) for s in slide.shapes if s.has_text_frame) for slide in prs.slides]
    for index, chart in zip([11, 16, 17], charts):
        counts[index] += chart['visible_words_upper_bound']
    if max(counts) > 80 or sum(counts) > 1100:
        raise ValueError(f'word budget exceeded: {counts}, total {sum(counts)}')
    manifest = {'p4_commit': sha, 'base_deck_sha256': hashlib.sha256(original).hexdigest(),
                'status': 'DRAFT: P1 approval and release J1 pending', 'slide_words_upper_bound': counts,
                'total_words_upper_bound': sum(counts), 'source_sha256': {p: hashlib.sha256(b).hexdigest() for p,b in blobs.items()},
                'artifacts': {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in output.iterdir() if p.is_file() and p.name != 'manifest.json'}}
    (output / 'manifest.json').write_text(json.dumps(manifest, indent=2)+'\n')
    print(f'Draft built: {len(prs.slides)} slides; conservative word count {sum(counts)}, max {max(counts)}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', type=Path, default=ROOT/'docs/presentation/draft_2026-10-16')
    build(parser.parse_args().out)
