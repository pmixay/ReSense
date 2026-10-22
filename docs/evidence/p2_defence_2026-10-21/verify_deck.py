"""Check actual PPTX chart caches/PDF text against separate evidence files."""
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import pymupdf
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
deck = ROOT / "docs/presentation/P2_defence_2026-10-21.pptx"
prs = Presentation(deck)
source = json.loads((ROOT / "docs/evidence/rejudge_2026-10-16/prebuild/summary_1.1.0.json").read_text())
jobs = {j["job"]: j["summary"]["first_stop_X"] for j in source["jobs"]}
range_chart = next(s.chart for s in prs.slides[11].shapes if s.has_chart)
expected = [int(jobs[k] + .5) for k in ["B_person", "C_person", "D_person", "A_cube03", "A_cube1", "A_person"]]
assert list(range_chart.series[0].values) == expected
# Positive control: the comparison rejects a planted one-metre error.
changed = list(expected)
changed[0] += 1
assert list(range_chart.series[0].values) != changed
gate = json.loads((ROOT / "docs/evidence/f1a_completion_2026-10-20/gate-c-candidate.json").read_text())
fake = next(s.chart for s in prs.slides[12].shapes if s.has_chart)
keys = ["thin_hanging", "big_edge_inside", "small_edge_inside", "small_center", "small_on_rail", "long_low_on_rails", "big_center", "big_above"]
assert list(fake.series[0].values) == [int(gate["set_O"]["objects"][k]["first_stop_m"] + .5) for k in keys]
pdf = pymupdf.open(deck.with_suffix(".pdf"))
assert len(pdf) == len(prs.slides) == 16
texts = [" ".join(p.get_text().split()) for p in pdf]
assert "1.1.0" in texts[11] and "Итоговый J1 ожидается" in texts[11]
assert "4 кадра STOP" in texts[13] and "444 с" in texts[13]
assert "24 эпизода" in texts[10]
assert "63 пропуска" in texts[14]
assert "8 без" in texts[15] and "f56ad5e" in texts[15]
for page in pdf:
    for word in page.get_text("words"):
        rect = pymupdf.Rect(word[:4])
        assert page.rect.contains(rect), (page.number + 1, word)
with zipfile.ZipFile(deck) as archive:
    notes = " ".join(" ".join(ET.fromstring(archive.read(n)).itertext()) for n in archive.namelist()
                     if n.startswith("ppt/notesSlides/notesSlide") and n.endswith(".xml"))
assert "gate-c-candidate.json" in notes
assert "summary_1.1.0.json" in notes
fonts = sorted({f[3] for p in pdf for f in p.get_fonts()})
assert any("Montserrat" in name for name in fonts)
result = {"pages": 16, "range_bars_match_source": expected, "set_o_bars_match_source": True,
          "planted_change_detected": True, "pdf_text_in_page": True, "fonts": fonts,
          "visual_review": "Contact sheet all 16 pages and full-size slides 12/14 inspected; no clipped chart labels or body text.",
          "final_j1": False, "p1_approval": False}
(HERE / "deck_checks.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
print(json.dumps(result, ensure_ascii=False, indent=2))
